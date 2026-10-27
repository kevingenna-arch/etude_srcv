"""Version Word de la page partagée « Bilan logement et natalité ».

Usage : python construire_bilan_word.py  (depuis ce dossier)
Lit bilan_logement_natalite.html, source de la page publiée, et écrit
bilan_logement_natalite.docx dans le style de la synthèse Word
(../synthese/construire_synthese.py). Les deux graphiques en barres viennent
de figures_bilan.R ; légendes et graduations HTML sont alors inutiles.
"""
import re
import sys
from html.parser import HTMLParser
from pathlib import Path

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ICI = Path(__file__).resolve().parent
sys.path.insert(0, str(ICI.parent / "synthese"))
from construire_synthese import (FOND, GRIS, MARINE, POLICE, Synthese,  # noqa: E402
                                 bordures_horizontales, marges_cellule, ombrer)

SOURCE = ICI / "bilan_logement_natalite.html"
SORTIE = ICI / "bilan_logement_natalite.docx"
FIGURES = [ICI / "graph_parts_icf_quartiles.png", ICI / "graph_parts_icf_diplome.png"]
OCRE = RGBColor(0x8F, 0x5A, 0x07)
PUCES = {"ok": RGBColor(0x2C, 0x7A, 0x4B), "part": RGBColor(0x94, 0x60, 0x06),
         "no": RGBColor(0xAD, 0x3A, 0x3A)}
LARGEUR = 16.2  # cm, entre les marges


# ── Lecture de la page ──────────────────────────────────────────────────────
class Noeud:
    def __init__(self, tag, attrs, parent):
        self.tag, self.attrs, self.parent = tag, dict(attrs), parent
        self.enfants = []

    @property
    def classes(self):
        return (self.attrs.get("class") or "").split()

    def elements(self, tag=None):
        return [e for e in self.enfants if isinstance(e, Noeud) and (tag is None or e.tag == tag)]

    def texte(self):
        return "".join(e if isinstance(e, str) else e.texte() for e in self.enfants)


class Arbre(HTMLParser):
    VIDES = {"meta", "link", "br", "img", "hr", "input"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.racine = Noeud("racine", [], None)
        self.courant = self.racine

    def handle_starttag(self, tag, attrs):
        n = Noeud(tag, attrs, self.courant)
        self.courant.enfants.append(n)
        if tag not in self.VIDES:
            self.courant = n

    def handle_endtag(self, tag):
        n = self.courant
        while n is not None and n.tag != tag:
            n = n.parent
        if n is not None:
            self.courant = n.parent

    def handle_data(self, data):
        self.courant.enfants.append(data)


def chercher(noeud, tag, classe=None):
    for e in noeud.elements():
        if e.tag == tag and (classe is None or classe in e.classes):
            return e
        trouve = chercher(e, tag, classe)
        if trouve is not None:
            return trouve
    return None


# ── Texte en ligne ──────────────────────────────────────────────────────────
def segments(noeud, gras=False, italique=False, exposant=False, couleur=None):
    """Morceaux de texte d'un nœud, avec leur mise en forme."""
    for e in noeud.enfants:
        if isinstance(e, str):
            # Espaces ordinaires seulement : les insécables servent de retrait
            yield (re.sub(r"[ \t\r\n]+", " ", e), gras, italique, exposant, couleur)
        elif e.tag in ("strong", "b"):
            yield from segments(e, True, italique, exposant, couleur)
        elif e.tag in ("em", "i"):
            yield from segments(e, gras, True, exposant, couleur)
        elif e.tag == "sup":
            yield from segments(e, gras, italique, True, couleur)
        elif "chip" in e.classes:
            teinte = next((PUCES[c] for c in e.classes if c in PUCES), couleur)
            yield from segments(e, True, italique, exposant, teinte)
        else:
            yield from segments(e, gras, italique, exposant, couleur)


def ecrire(par, morceaux, taille=None, couleur=None, gras=False, italique=False):
    """Écrit les morceaux dans le paragraphe, espaces insécables compris."""
    morceaux = [list(m) for m in morceaux if m[0]]
    if morceaux:
        morceaux[0][0] = morceaux[0][0].lstrip(" ")
        morceaux[-1][0] = morceaux[-1][0].rstrip(" ")
    for t, g, i, sup, c in morceaux:
        if not t:
            continue
        t = re.sub(r"(\d) (\d{3})\b", "\\1\u00a0\\2", t)
        t = t.replace(" %", "\u00a0%").replace(" €", "\u00a0€").replace(" :", "\u00a0:")
        r = par.add_run(t)
        r.bold = gras or g
        r.italic = italique or i
        r.font.superscript = sup
        if taille:
            r.font.size = Pt(taille)
        if (c or couleur) is not None:
            r.font.color.rgb = c or couleur
    return par


# ── Document ────────────────────────────────────────────────────────────────
class Bilan(Synthese):
    def __init__(self):
        super().__init__()
        styles = self.doc.styles
        rpr = styles["Normal"].element.get_or_add_rPr()
        lang = OxmlElement("w:lang")
        lang.set(qn("w:val"), "fr-FR")
        rpr.append(lang)
        h = styles["Heading 3"]
        rf = h.element.rPr.rFonts
        for att in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            rf.attrib.pop(qn(f"w:{att}"), None)
        for att in ("ascii", "hAnsi", "eastAsia", "cs"):
            rf.set(qn(f"w:{att}"), POLICE)
        h.font.size = Pt(11.5)
        h.font.bold, h.font.italic = True, True
        h.font.color.rgb = MARINE
        h.paragraph_format.space_before = Pt(8)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
        self.figures = iter(FIGURES)

    def titre_page(self, header):
        sur = chercher(header, "div", "eyebrow")
        p = self.doc.add_paragraph()
        ecrire(p, [(sur.texte().upper(), True, False, False, None)], taille=9, couleur=OCRE)
        p.paragraph_format.space_after = Pt(4)
        p = self.doc.add_paragraph()
        ecrire(p, segments(chercher(header, "h1")), taille=24, couleur=MARINE)
        p.paragraph_format.space_after = Pt(2)
        p = self.doc.add_paragraph()
        ecrire(p, segments(chercher(header, "p", "standfirst")), taille=13, italique=True,
               couleur=MARINE)
        p = self.doc.add_paragraph()
        ecrire(p, segments(chercher(header, "p", "meta")), taille=10, couleur=GRIS)
        p.paragraph_format.space_after = Pt(12)

    def encadre_page(self, div):
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = t.cell(0, 0)
        ombrer(c, FOND)
        marges_cellule(c, 0.35)
        ecrire(c.paragraphs[0], segments(chercher(div, "div", "label")), gras=True, taille=12,
               couleur=MARINE)
        for li in chercher(div, "ul").elements("li"):
            par = c.add_paragraph(style="List Bullet")
            par.paragraph_format.space_after = Pt(3)
            ecrire(par, segments(li))
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def puces_page(self, ul):
        for li in ul.elements("li"):
            par = self.doc.add_paragraph(style="List Bullet")
            par.paragraph_format.space_after = Pt(3)
            ecrire(par, segments(li))

    def paragraphe(self, p):
        par = self.doc.add_paragraph()
        if "note" in p.classes:
            ecrire(par, segments(p), taille=9.5, italique=True, couleur=GRIS)
        else:
            ecrire(par, segments(p))

    def figure(self):
        self.doc.add_picture(str(next(self.figures)), width=Cm(16))
        self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    def tableau_page(self, table):
        lignes = [tr.elements() for tr in chercher(table, "thead").elements("tr")]
        lignes += [tr.elements() for tr in chercher(table, "tbody").elements("tr")]
        ncol = sum(int(c.attrs.get("colspan", 1)) for c in lignes[0])
        taille = 10 if ncol <= 4 else 9 if ncol <= 6 else 8.5
        largeurs = largeurs_colonnes(lignes, ncol, taille)

        t = self.doc.add_table(rows=len(lignes), cols=ncol)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        bordures_horizontales(t)
        # Largeurs avant fusion : une cellule fusionnée prend la somme des siennes
        for col, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), largeurs):
            col.set(qn("w:w"), str(int(w * 567)))
        for rangee in t.rows:
            for j, w in enumerate(largeurs):
                rangee.cells[j].width = Cm(w)
        for i, cells in enumerate(lignes):
            rangee = t.rows[i]
            trPr = rangee._tr.get_or_add_trPr()
            trPr.append(OxmlElement("w:cantSplit"))
            if i == 0:
                trPr.append(OxmlElement("w:tblHeader"))
            j = 0
            for c in cells:
                n = int(c.attrs.get("colspan", 1))
                cellule = rangee.cells[j]
                if n > 1:
                    cellule = cellule.merge(rangee.cells[j + n - 1])
                par = cellule.paragraphs[0]
                if c.tag == "th":
                    ombrer(cellule, FOND)
                    ecrire(par, segments(c), gras=True, taille=taille, couleur=MARINE)
                elif "doc" in c.classes:
                    sous = chercher(c, "span")
                    titre = [e for e in c.enfants if e is not sous]
                    ecrire(par, segments(noeud_de(titre)), gras=True, taille=taille)
                    ecrire(cellule.add_paragraph(), segments(sous), taille=taille - 1,
                           couleur=GRIS)
                else:
                    ecrire(par, segments(c), taille=taille)
                    if c.texte().startswith("  "):  # sous-ligne (« dont … »)
                        par.paragraph_format.left_indent = Cm(0.4)
                if "num" in c.classes:
                    for p in cellule.paragraphs:
                        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                j += n
        # Tableau d'un seul tenant : chaque rangée reste avec la suivante
        for i, rangee in enumerate(t.rows):
            for cellule in rangee.cells:
                for p in cellule.paragraphs:
                    p.paragraph_format.space_after = Pt(1)
                    p.paragraph_format.line_spacing = 1.0
                    p.paragraph_format.keep_with_next = i < len(t.rows) - 1
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # Parcours
    def bloc(self, n, niveau):
        """Rend un bloc ; niveau = niveau de titre des h3 rencontrés."""
        if not isinstance(n, Noeud):
            return
        cl = n.classes
        if n.tag == "h3":
            # « 3.5 … » est un titre de niveau 2 ; « Par quartile … » en dépend
            numerote = re.match(r"\d", n.texte().strip())
            self.doc.add_heading(n.texte().strip(), level=2 if niveau == 2 and numerote else 3)
        elif n.tag == "p":
            self.paragraphe(n)
        elif n.tag == "ul":
            self.puces_page(n)
        elif "table-wrap" in cl:
            self.tableau_page(chercher(n, "table"))
        elif "barres" in cl:
            self.figure()
        elif "legende" in cl or "graduation" in cl:
            return
        elif "brief" in cl:
            self.encadre_page(n)
        elif "cols" in cl:
            # Colonnes directement sous une section (partie 4) : titres de niveau 2
            sous_section = n.parent.tag == "section"
            for col in n.elements("div"):
                for e in col.elements():
                    if e.tag == "h3":
                        self.doc.add_heading(e.texte().strip(), level=2 if sous_section else 3)
                    else:
                        self.bloc(e, 3)
        else:
            for e in n.elements():
                self.bloc(e, niveau)

    def section(self, sec):
        for e in sec.elements():
            if e.tag == "h2":
                self.doc.add_heading(e.texte().strip(), level=1)
            else:
                self.bloc(e, 2)


def largeurs_colonnes(lignes, ncol, taille):
    """Largeurs en cm : les colonnes de chiffres tiennent sur une ligne, aucune
    colonne ne coupe un mot ; le reste va aux colonnes de texte, au prorata de
    leur texte le plus long."""
    car = 0.19 * taille / 10  # largeur moyenne d'un caractère, en cm
    marge = 0.45
    longueur = [0] * ncol   # texte le plus long (plafonné) de la colonne
    mot = [4] * ncol        # mot le plus long
    nombre = [0] * ncol     # cellule de chiffres la plus longue
    entete = [0] * ncol     # longueur de l'en-tête
    chiffres = [True] * ncol
    for i, cells in enumerate(lignes):
        j = 0
        for c in cells:
            n = int(c.attrs.get("colspan", 1))
            if n == 1:
                t = re.sub(r"\s+", " ", c.texte()).strip()
                longueur[j] = max(longueur[j], min(len(t), 60))
                mot[j] = max([mot[j]] + [len(m) for m in t.split()])
                if c.tag == "th":
                    entete[j] = max(entete[j], len(t))
                if i > 0 and c.tag == "td":
                    chiffres[j] &= "num" in c.classes
                    if "num" in c.classes:
                        nombre[j] = max(nombre[j], len(t))
            j += n
    # Chiffres : une ligne jusqu'à 11 caractères, au-delà coupure entre les mots ;
    # en-tête sur deux lignes au plus
    mini = [(max(min(nombre[j], 11), min(entete[j], 10), mot[j]) if chiffres[j] else mot[j])
            * car + marge for j in range(ncol)]
    texte_ = [j for j in range(ncol) if not chiffres[j]]
    poids = {j: max(longueur[j], 24 if j == 0 else 8) for j in texte_}
    reste = LARGEUR - sum(mini)
    largeurs = list(mini)
    for j in texte_:
        largeurs[j] += max(reste, 0) * poids[j] / sum(poids.values())
    if not texte_ or reste < 0:
        largeurs = [w * LARGEUR / sum(largeurs) for w in largeurs]
    return largeurs


def noeud_de(enfants):
    n = Noeud("span", [], None)
    n.enfants = enfants
    return n


def main():
    arbre = Arbre()
    arbre.feed(SOURCE.read_text(encoding="utf-8"))
    page = chercher(arbre.racine, "div", "page")
    s = Bilan()
    for e in page.elements():
        if e.tag == "header":
            s.titre_page(e)
        elif "brief" in e.classes:
            s.encadre_page(e)
        elif e.tag == "section":
            s.section(e)
        elif e.tag == "footer":
            s.doc.add_heading("Sources", level=2)
            texte_sources = chercher(e, "p").texte().strip()
            texte_sources = re.sub(r"^Sources\s*:\s*", "", texte_sources)
            par = s.doc.add_paragraph()
            ecrire(par, [(texte_sources[0].upper() + texte_sources[1:], False, False, False, None)],
                   taille=10, couleur=GRIS)
    s.doc.core_properties.title = "Bilan logement et natalité"
    s.doc.core_properties.author = "Kévin Genna"
    s.doc.save(SORTIE)
    print("Écrit :", SORTIE)


if __name__ == "__main__":
    main()
