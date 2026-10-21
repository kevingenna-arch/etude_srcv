"""Document Word de synthèse : logement, politique familiale et baisse de la fécondité.

Usage : python construire_synthese.py  (depuis ce dossier)
Produit synthese_logement_fecondite.docx. Les chiffres de la politique
familiale sont lus dans les sorties du script R/38 (Output/, hors dépôt).
"""
import csv
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ICI = Path(__file__).resolve().parent
OUTPUT = ICI.parent.parent / "Output"
SORTIE = ICI / "synthese_logement_fecondite.docx"

MARINE = RGBColor(0x24, 0x34, 0x43)
GRIS = RGBColor(0x59, 0x59, 0x59)
FOND = "EEF1F4"
POLICE = "Garamond"


def lire(nom):
    with open(OUTPUT / nom, encoding="utf-8") as f:
        return list(csv.DictReader(f))


def fr(x, d=1):
    """Nombre au format français."""
    s = f"{float(x):,.{d}f}".replace(",", " ").replace(".", ",")
    return s.replace("-", "−")


# ── Mise en forme ───────────────────────────────────────────────────────────
def ombrer(cellule, couleur):
    tc = cellule._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), couleur)
    tc.append(shd)


def marges_cellule(cellule, cm=0.15):
    tcPr = cellule._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for cote in ("top", "bottom", "left", "right"):
        e = OxmlElement(f"w:{cote}")
        e.set(qn("w:w"), str(int(cm * 567)))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tcPr.append(mar)


def bordures_horizontales(table):
    """Filets horizontaux seulement, à la manière de booktabs."""
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for cote, taille in (("top", 12), ("bottom", 12), ("insideH", 4)):
        e = OxmlElement(f"w:{cote}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(taille))
        e.set(qn("w:color"), "8A949E" if cote == "insideH" else "243443")
        b.append(e)
    for cote in ("left", "right", "insideV"):
        e = OxmlElement(f"w:{cote}")
        e.set(qn("w:val"), "nil")
        b.append(e)
    tblPr.append(b)


def texte(par, contenu, gras=False, italique=False, taille=None, couleur=None):
    """Ajoute du texte ; **gras** dans la chaîne est interprété."""
    # Espaces insécables : avant % et €, et dans les milliers (26 470)
    contenu = re.sub(r"(\d) (\d{3})\b", "\\1\u00a0\\2", contenu)
    contenu = contenu.replace(" %", "\u00a0%").replace(" €", "\u00a0€")
    morceaux = contenu.split("**")
    for i, m in enumerate(morceaux):
        if not m:
            continue
        r = par.add_run(m)
        r.bold = gras or (i % 2 == 1)
        r.italic = italique
        if taille:
            r.font.size = Pt(taille)
        if couleur is not None:
            r.font.color.rgb = couleur
    return par


class Synthese:
    def __init__(self):
        self.doc = Document()
        s = self.doc.sections[0]
        s.page_height, s.page_width = Cm(29.7), Cm(21)
        s.top_margin = s.bottom_margin = Cm(2.2)
        s.left_margin = s.right_margin = Cm(2.4)
        st = self.doc.styles["Normal"]
        st.font.name = POLICE
        st.element.rPr.rFonts.set(qn("w:eastAsia"), POLICE)
        st.font.size = Pt(11.5)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.line_spacing = 1.12
        for niveau, taille in ((1, 16), (2, 13)):
            h = self.doc.styles[f"Heading {niveau}"]
            h.font.name = POLICE
            # Les titres héritent des polices du thème, que Word applique encore aux
            # lettres accentuées : on les retire et on fixe les quatre familles.
            rf = h.element.rPr.rFonts
            for att in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
                rf.attrib.pop(qn(f"w:{att}"), None)
            for att in ("ascii", "hAnsi", "eastAsia", "cs"):
                rf.set(qn(f"w:{att}"), POLICE)
            h.font.size = Pt(taille)
            h.font.bold = niveau == 2
            h.font.color.rgb = MARINE
            h.paragraph_format.space_before = Pt(16 if niveau == 1 else 10)
            h.paragraph_format.space_after = Pt(6)
            h.paragraph_format.keep_with_next = True
        for nom in ("List Bullet", "List Number"):
            self.doc.styles[nom].font.name = POLICE
        pied = s.footer.paragraphs[0]
        pied.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._numero_page(pied)

    def _numero_page(self, par):
        r = par.add_run()
        for typ, txt in (("begin", None), (None, "PAGE"), ("end", None)):
            if typ:
                e = OxmlElement("w:fldChar")
                e.set(qn("w:fldCharType"), typ)
            else:
                e = OxmlElement("w:instrText")
                e.set(qn("xml:space"), "preserve")
                e.text = txt
            r._r.append(e)
        r.font.size = Pt(9)
        r.font.color.rgb = GRIS

    # Blocs
    def titre(self, t, sous_titre, auteur, date):
        p = self.doc.add_paragraph()
        texte(p, t, taille=22, couleur=MARINE)
        p.paragraph_format.space_after = Pt(2)
        p = self.doc.add_paragraph()
        texte(p, sous_titre, italique=True, taille=13, couleur=MARINE)
        p = self.doc.add_paragraph()
        texte(p, f"{auteur} · {date}", taille=10, couleur=GRIS)
        p.paragraph_format.space_after = Pt(12)

    def h1(self, t):
        self.doc.add_heading(t, level=1)

    def h2(self, t):
        self.doc.add_heading(t, level=2)

    def p(self, t, **kw):
        return texte(self.doc.add_paragraph(), t, **kw)

    def puces(self, items):
        for it in items:
            par = self.doc.add_paragraph(style="List Bullet")
            par.paragraph_format.space_after = Pt(3)
            texte(par, it)

    def encadre(self, titre, items):
        t = self.doc.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = t.cell(0, 0)
        ombrer(c, FOND)
        marges_cellule(c, 0.35)
        par = c.paragraphs[0]
        texte(par, titre, gras=True, taille=12, couleur=MARINE)
        for it in items:
            par = c.add_paragraph(style="List Bullet")
            par.paragraph_format.space_after = Pt(3)
            texte(par, it)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def tableau(self, entetes, lignes, largeurs, note=None, taille=10):
        t = self.doc.add_table(rows=1, cols=len(entetes))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        bordures_horizontales(t)
        for j, e in enumerate(entetes):
            c = t.rows[0].cells[j]
            ombrer(c, FOND)
            texte(c.paragraphs[0], e, gras=True, taille=taille, couleur=MARINE)
        for ligne in lignes:
            cells = t.add_row().cells
            for j, v in enumerate(ligne):
                texte(cells[j].paragraphs[0], str(v), taille=taille)
        for row in t.rows:
            for j, w in enumerate(largeurs):
                row.cells[j].width = Cm(w)
                for par in row.cells[j].paragraphs:
                    par.paragraph_format.space_after = Pt(1)
                    par.paragraph_format.line_spacing = 1.0
        if note:
            par = self.doc.add_paragraph()
            par.paragraph_format.space_before = Pt(3)
            texte(par, note, italique=True, taille=9, couleur=GRIS)
        else:
            self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def enregistrer(self):
        self.doc.save(SORTIE)
        print("Écrit :", SORTIE)



# ── Contenu (période 2014-2025) ─────────────────────────────────────────────
def en_bref(s, pf):
    s.encadre("En bref", [
        "L'ICF a baissé de 22 % entre 2014 et 2025 (de 2,00 à 1,56). Aucun facteur n'explique "
        "à lui seul cette baisse ; ceux que nous avons mesurés en couvrent environ les deux "
        "tiers.",
        "Le premier facteur identifié est le **recul de la vie en couple** : environ un tiers de "
        "la baisse, car 92 % des naissances ont lieu en couple.",
        "Les **petites surfaces** pèsent de 4 à 19 % de la baisse selon la spécification. Mais "
        "le recul des propriétaires non accédants a joué en sens inverse, et le logement dans "
        "son ensemble n'explique presque rien ; l'association ressemble d'ailleurs davantage à "
        "un choix qu'à une contrainte.",
        f"Le barème de la **politique familiale** est devenu moins généreux : en euros constants, "
        f"un enfant de plus rapporte {fr(pf['perte'], 0)} € de moins par an qu'avec le barème de "
        f"2014 ({fr(-pf['perte_pct'], 0)} %), pour les trois quarts à cause de la PAJE. Ce recul "
        f"pourrait expliquer de l'ordre de {fr(pf['part14'], 0)} % de la baisse, mais "
        "l'estimation est imprécise et ne permet pas de l'affirmer.",
        "Un tiers de la baisse reste inexpliqué : à situation donnée, les couples font moins "
        "d'enfants, ce qui renvoie au désir d'enfant (en net recul selon l'Ined) et à la "
        "confiance en l'avenir.",
    ])


def contexte(s):
    s.h1("1. Le point de départ")
    s.p("L'indicateur conjoncturel de fécondité (ICF) est passé de 2,00 enfants par femme en "
        "2014 à 1,56 en 2025, son niveau le plus bas depuis la fin de la Première Guerre "
        "mondiale, soit une baisse de 22 %. En 2025, les décès ont dépassé les naissances pour "
        "la première fois depuis la fin de la Seconde Guerre mondiale (Insee Première n° 2087).")
    s.p("Nous avons cherché à mesurer ce que le logement, puis la politique familiale, "
        "expliquent de cette baisse, à partir de l'enquête Statistiques sur les ressources et "
        "conditions de vie (SRCV) de l'Insee : six vagues, 2014, 2018 puis chaque année de 2022 "
        "à 2025. Le champ principal compte 19 645 couples dont la femme a 15 à 49 ans, et 1 519 "
        "naissances.")
    s.p("La méthode est la même pour chaque facteur. Un modèle estime le lien entre ce facteur "
        "et la probabilité d'une naissance, toutes choses égales par ailleurs ; on donne ensuite "
        "aux couples de 2025 la valeur que ce facteur avait en 2014, et l'on mesure de combien "
        "leur natalité aurait été plus forte. Rapporté à la baisse de l'ICF, cet écart donne la "
        "« part de la baisse » imputable au facteur. Ces parts sont des bornes hautes : elles "
        "lisent comme causale une association qui ne l'est pas forcément.")


def vue_ensemble(s, pf, diplomes):
    s.h1("2. Ce qui explique la baisse : vue d'ensemble")
    s.p("La figure 1 rassemble la part de la baisse de l'ICF entre 2014 et 2025 associée à "
        "chacun des facteurs étudiés. Les valeurs sont centrales ; la fourchette figure sous "
        "chaque segment. La politique familiale est en teinte claire parce que son estimation "
        "est imprécise.")
    s.doc.add_picture(str(ICI / "graph_parts_baisse_icf.png"), width=Cm(16))
    s.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    par = s.doc.add_paragraph()
    texte(par, "Figure 1 — Part de la baisse de l'ICF entre 2014 et 2025 associée à chaque "
          "facteur. ", gras=True, taille=9.5)
    texte(par, "Ne figure pas dans la barre : la structure des statuts d'occupation, qui a "
          "soutenu la natalité (environ −10 %), de sorte que le logement dans son ensemble pèse "
          "entre −6 % et +8 %. Les parts ne s'additionnent qu'approximativement : elles viennent "
          "de méthodes et de champs différents.", italique=True, taille=9.5, couleur=GRIS)
    s.tableau(
        ["Facteur", "Part de la baisse", "Fourchette", "Lecture"],
        [["Recul de la vie en couple", "34 %", "33 à 34 %",
          "Moins de femmes en couple à âge donné ; corrélation, le lien avec le logement n'est "
          "pas établi"],
         ["Logement : petites surfaces", "12 %", "4 à 19 %",
          "Effet de la surface significatif, mais probablement un choix plus qu'une contrainte"],
         ["Logement : statuts d'occupation", "−10 %", "",
          "Moins de propriétaires non accédants, les moins féconds : la structure a soutenu la "
          "natalité"],
         ["Politique familiale (barème)", f"{fr(pf['part14'], 0)} %",
          f"{fr(pf['bas14'], 0)} à {fr(pf['haut14'], 0)} %",
          "Non significatif ; tient presque entièrement à la PAJE ; porte sur les naissances de "
          "rang 2 ou plus"],
         ["Âge des couples", "4 %", "2 à 4 %", "Les couples ont vieilli d'un demi-an"],
         ["Prix au m², revenu, concentration urbaine", "≈ 0", "",
          "Aucun effet à surface donnée ; la natalité n'est pas plus faible en ville"],
         ["Non expliqué", "33 %", "", "Moins de naissances à situation donnée"]],
        [4.3, 2.2, 2.3, 7.4],
        note="Parts de la baisse de l'ICF entre 2014 et 2025 (−22 %), sous l'hypothèse que la "
             "natalité des femmes hors du champ étudié varie comme celle du champ. Logement et "
             "politique familiale : modèle sur les couples, logement prédéterminé ; seuil de "
             "surface à 60 m² pour la valeur centrale.")
    s.h2("Par quartile de niveau de vie")
    s.p("La figure 2 refait le même calcul dans chaque quartile de niveau de vie. Les "
        "quartiles portent sur le niveau de vie du ménage des femmes de 15 à 49 ans, hors "
        "prestations ; chaque facteur est mesuré avec la même méthode que pour l'ensemble, mais "
        "sur les seules femmes ou les seuls couples du quartile, à une exception près : le "
        "recul de l'âge des couples est mesuré au niveau national puis appliqué à chaque "
        "groupe (mesuré groupe par groupe, il reflétait surtout la recomposition des groupes "
        "entre 2014 et 2025). Toutes les barres sont rapportées à la baisse nationale "
        "(−22 %) : une contribution négative part à gauche de zéro.")
    s.doc.add_picture(str(ICI / "graph_parts_icf_quartiles.png"), width=Cm(16))
    s.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    par = s.doc.add_paragraph()
    texte(par, "Figure 2 — Part de la baisse de l'ICF entre 2014 et 2025 associée à chaque "
          "facteur, par quartile de niveau de vie. ", gras=True, taille=9.5)
    texte(par, "Parts de la baisse nationale de l'ICF ; le « non expliqué » suppose que chaque "
          "quartile a baissé comme la moyenne. Le trait pointillé marque 100 %.",
          italique=True, taille=9.5, couleur=GRIS)
    s.puces([
        "**La politique familiale se concentre sur les classes moyennes** : 33 % de la baisse "
        "en Q3 et 15 % en Q2, contre 7 à 11 % aux extrémités. C'est le résultat le plus net, "
        "et il rejoint la section 4.4 : ce sont ces couples qui ont perdu la PAJE.",
        "**Le recul de la vie en couple atteint 66 % en Q1**, contre 26 à 33 % ailleurs, mais "
        "c'est surtout un effet de classement. Une femme qui ne vit pas en couple n'a qu'un "
        "revenu dans son ménage et tombe en Q1 : c'est la vie hors couple qui classe en Q1, "
        "et non un revenu modeste qui empêche de se mettre en couple. Comme les femmes hors "
        "couple sont plus nombreuses, Q1 s'en remplit (part de femmes en couple de 44 % à 35 %, "
        "contre 70 % à 69 % en Q4). Le résultat est le même avec des seuils de quartile fixés "
        "en 2014, et le découpage par diplôme, ci-dessous, ne montre pas de recul concentré "
        "chez les moins diplômées.",
        "**Le logement** va de 3 % (Q3) à 23 % (Q4), sans profil net selon le revenu ; "
        "**l'âge des couples** pèse peu partout (0 à 10 %).",
    ])
    s.h2("Par diplôme de la femme")
    s.p("Le diplôme a l'avantage de ne pas dépendre de la vie en couple. La figure 3 refait le "
        "calcul par diplôme, sur les femmes de 25 à 49 ans : avant 25 ans, beaucoup sont "
        "encore étudiantes, et leur diplôme du moment les classerait à tort parmi les moins "
        "diplômées.")
    s.doc.add_picture(str(ICI / "graph_parts_icf_diplome.png"), width=Cm(16))
    s.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    par = s.doc.add_paragraph()
    texte(par, "Figure 3 — Part de la baisse de l'ICF entre 2014 et 2025 associée à chaque "
          "facteur, par diplôme de la femme (25-49 ans). ", gras=True, taille=9.5)
    texte(par, "Parts de la baisse nationale de l'ICF ; le « non expliqué » suppose que chaque "
          "groupe a baissé comme la moyenne. Le trait pointillé marque 100 %.",
          italique=True, taille=9.5, couleur=GRIS)
    s.tableau(
        ["Diplôme de la femme (25-49 ans)", "En couple en 2014", "En couple en 2024-2025",
         "Part de la baisse liée à la vie en couple"],
        [[r["groupe"].replace("Ensemble", "Ensemble des 25-49 ans"),
          f"{fr(r['en_couple_2014'], 0)} %", f"{fr(r['en_couple_2425'], 0)} %",
          f"{fr(r['vie_en_couple'], 0)} %"] for r in diplomes],
        [6.2, 3.0, 3.2, 3.8], taille=9.5,
        note="Femmes de 15 à 49 ans pour le découpage par quartile, de 25 à 49 ans pour le "
             "découpage par diplôme : les deux lignes « Ensemble » diffèrent pour cette raison.")
    s.puces([
        "**Le recul de la vie en couple touche tous les niveaux de diplôme** : il pèse 25 à "
        "29 % de la baisse, sauf chez les bac+2 (50 %, où la part en couple passe de 79 % à "
        "72 %). Il n'est pas concentré chez les femmes peu diplômées.",
        "**La politique familiale pèse autant dans tous les groupes** (17 à 20 %) : le barème "
        "dépend du revenu, pas du diplôme, et les deux ne se recouvrent qu'en partie.",
        "**Le logement** pèse surtout aux deux extrémités (18 % pour le supérieur, 16 % pour "
        "CAP-BEP ou moins) et presque rien pour le bac et le bac+2 ; avec 34 à 93 naissances "
        "par groupe en 2025, ces écarts restent fragiles.",
    ])


def logement(s):
    s.h1("3. Le logement")
    s.h2("3.1 Les conditions de logement des jeunes")
    s.puces([
        "Entre 2014 et 2025, les locataires de 20 à 39 ans ont perdu de 1 à 6 m² selon l'âge "
        "(de −2 % à −8 %), alors que leur taux d'effort reculait de 1,4 à 4,8 points : ils ont "
        "perdu de l'espace, pas du pouvoir d'achat.",
        "Les propriétaires accédants ont gardé leurs surfaces et vu leur taux d'effort baisser "
        "de 4 à 8 points, grâce aux taux d'intérêt bas.",
        "La surface moyenne des couples en âge d'avoir des enfants a peu varié (de 102,8 à "
        "100,9 m²), mais la médiane a perdu 6 m², et la part des couples sans enfant logés sous "
        "50 m² est passée de 14,6 % à 19,1 %.",
    ])
    s.h2("3.2 Surface et naissance")
    s.puces([
        "Dix mètres carrés de plus sont associés à une probabilité de naissance plus élevée de "
        "0,34 point (environ +4 %), et de 0,59 point pour le premier enfant.",
        "Le **prix au m²** n'a aucun effet à surface donnée, et le **revenu** non plus, à condition "
        "de le mesurer avant la naissance et hors prestations.",
        "Pour le premier enfant, la relation prend la forme d'un **seuil** : sous 50 m², la "
        "probabilité d'une première naissance est inférieure de 9 points (sur un taux moyen de "
        "13,8 %). Le seuil se situe entre 45 et 60 m² ; au-dessus, la surface ne départage plus "
        "rien.",
        "Pour le premier enfant, les accédants ont un avantage de 5 points sur les locataires, "
        "que ni l'âge, ni le revenu, ni la catégorie sociale n'expliquent.",
        "Ces résultats résistent à l'âge du conjoint, au type de logement, à un indice de prix "
        "local (DVF), au double comptage des naissances et à la mesure de l'urbanisation.",
    ])
    s.h2("3.3 Contrainte ou choix ?")
    s.p("Le panel suit les mêmes couples d'une vague à l'autre. La surface y est liée à la "
        "naissance récente (+0,26 point) mais ne prédit pas significativement la suivante "
        "(+0,09 point, intervalle de −0,14 à +0,31). L'association ressemble donc davantage à un "
        "logement choisi en vue d'un projet d'enfant qu'à une contrainte qui empêcherait les "
        "naissances, sans que le panel, trop court, puisse exclure un effet causal. La "
        "géographie va dans le même sens : les logements urbains sont plus petits, mais la "
        "fécondité n'est pas plus faible en ville.")
    s.h2("3.4 Le modèle théorique")
    s.p("Un modèle à générations imbriquées formalise la contrainte d'espace : le jeune adulte "
        "achète un logement, puis décide de sa fécondité dans ce logement. Ses prédictions en "
        "coupe se vérifient (la surface compte, pas son prix ; le revenu n'agit qu'à travers la "
        "surface), mais pas la prédiction centrale d'ordre temporel. Calibré, le canal d'espace "
        "pur est faible (élasticité de 0,11, contre 0,4 dans les données), ce qui plaide aussi "
        "pour la lecture « choix ».")


def politique_familiale(s, pf, gains, instr, quart):
    s.h1("4. La politique familiale")
    s.h2("4.1 Ce que montrent les prestations reçues")
    s.puces([
        "Pour les couples qui viennent d'avoir un enfant, le poids des prestations familiales "
        "dans le revenu a reculé entre 2014 et 2025 : de 6,7 % à 3,8 % au premier enfant, de "
        "10,4 % à 7,8 % au deuxième, de 14,7 % à 13,2 % au troisième ou plus.",
        "Le quotient familial rapporte 84 € par an au premier quintile de revenu, contre 1 928 € "
        "au dernier ; 86 % des ménages du premier quintile n'en tirent rien.",
        "Le maintien en emploi des mères de jeunes enfants dépend d'abord du diplôme (84 % en "
        "emploi pour les plus diplômées, 35 % pour les moins diplômées).",
    ])
    s.h2("4.2 Mesurer l'incitation, comme pour le logement")
    s.p("Les prestations reçues ne peuvent pas entrer telles quelles dans le modèle : c'est la "
        "naissance qui les déclenche. Nous avons donc calculé, pour chaque couple et avant la "
        "naissance, ce qu'un enfant de plus lui rapporterait selon le barème de l'année : "
        "allocations familiales, complément familial, allocation de base et prime de naissance "
        "de la PAJE, et gain de quotient familial. Le calcul utilise OpenFisca, le modèle libre "
        "du système socio-fiscal français, sur le revenu, l'activité des parents et l'âge des "
        "enfants déjà nés.")
    s.tableau(
        ["Gain d'un enfant supplémentaire, couples de 2025", "Barème 2014", "Barème 2025",
         "Écart"],
        [[lib, f"{fr(g[2014], 0)} €", f"{fr(g[2025], 0)} €",
          f"{fr(100 * (g[2025] / g[2014] - 1), 0)} %"] for lib, g in gains],
        [7.6, 2.8, 2.8, 2.8],
        note="Euros constants de 2025 par an, première année de l'enfant (prime de naissance "
             "comprise). Le barème de 2014 est appliqué aux mêmes couples, revenus ramenés en "
             "euros de 2014 : l'écart mesure ce qui va au-delà de l'indexation sur les prix. "
             "Couples supposés mariés.")
    s.h2("4.3 Pourquoi un tel recul ? Les trois instruments")
    s.p("La politique familiale passe par trois canaux : les prestations monétaires, les "
        "baisses d'impôt et les services en nature. Le tableau rend aux couples de 2025 un seul "
        "instrument du barème de 2014, les autres restant à leur niveau de 2025.")
    s.tableau(
        ["Instrument", "Barème 2014", "Barème 2025", "Écart", "Part de la baisse de l'ICF"],
        [[lib, f"{fr(r['gain_bareme'], 0)} €", f"{fr(r['gain_2025'], 0)} €",
          f"{fr(-r['ecart_gain'], 0)} €",
          f"{fr(r['part_de_la_baisse_pct'], 0)} % ({fr(r['part_ic_bas'], 0)} à "
          f"{fr(r['part_ic_haut'], 0)})"] for lib, r in instr]
        + [["Services en nature (crèches, assistantes maternelles)", "non mesuré", "non mesuré",
            "offre en hausse", "non chiffré"]],
        [6.1, 2.2, 2.2, 2.3, 3.4],
        note="Gain annuel d'un enfant supplémentaire, euros constants de 2025 ; part de la "
             "baisse de l'ICF 2014-2025 et intervalle de confiance à 95 %. Les sous-totaux "
             "« dont » s'ajoutent à la ligne des prestations monétaires.")
    s.puces([
        "**Les prestations monétaires portent presque tout le recul** (94 %), et d'abord la "
        "**PAJE** : son "
        "allocation de base a été réduite en 2014 (création d'un taux partiel pour les revenus "
        "intermédiaires), puis son montant et ses plafonds ont été alignés en 2018 sur ceux, "
        "plus bas, du complément familial. Pour les enfants de rang deux ou plus s'ajoute la "
        "modulation des allocations familiales selon le revenu, en 2015.",
        "**Les baisses d'impôt** n'ont guère bougé en moyenne : le plafond du quotient "
        "familial avait été abaissé en 2013 et 2014, avant notre période, et il est depuis "
        "indexé. Les écarts entre quartiles sont détaillés ci-dessous.",
        "**Les services en nature** ne sont pas observables dans SRCV et ne sont pas chiffrés "
        "ici. Ils n'ont en tout cas pas reculé : la capacité d'accueil des moins de 3 ans est "
        "passée d'environ 56 places pour 100 enfants en 2015 à 60,9 en 2023, en partie parce "
        "que les naissances baissent, et la dépense publique d'accueil du jeune enfant a "
        "continué de croître (17,5 milliards d'euros en 2024). Ce volet ne contribue donc pas à "
        "la baisse, même si Saudubray (2026) le juge le plus efficace à long terme.",
    ])
    s.h2("4.4 Qui a perdu ? Lecture par quartile de niveau de vie")
    s.p("Les réformes n'ont pas touché tous les couples de la même façon. La figure 4 et le "
        "tableau répartissent le recul du gain d'un enfant supplémentaire par quartile de "
        "niveau de vie, mesuré avant la naissance et hors prestations.")
    s.doc.add_picture(str(ICI / "graph_pf_quartiles.png"), width=Cm(14))
    s.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    par = s.doc.add_paragraph()
    texte(par, "Figure 4 — Recul du gain d'un enfant supplémentaire entre le barème de 2014 et "
          "celui de 2025, par quartile de niveau de vie. ", gras=True, taille=9.5)
    texte(par, "Couples de 2025, euros constants de 2025 par an. Sous chaque barre : l'écart "
          "total et l'effet estimé sur la natalité (ordre de grandeur, non significatif).",
          italique=True, taille=9.5, couleur=GRIS)
    s.tableau(
        ["Quartile", "Barème 2014", "Barème 2025", "Écart", "dont PAJE", "dont AF et CF",
         "dont quotient familial", "Effet estimé sur la natalité"],
        [[r["q"], f"{fr(r['gain_bareme_ref'], 0)} €", f"{fr(r['gain_2025'], 0)} €",
          f"{fr(r['ecart'], 0)} € ({fr(r['ecart_pct'], 0)} %)", f"{fr(r['ecart_paje'], 0)} €",
          f"{fr(r['ecart_af_cf'], 0)} €", f"{fr(r['ecart_qf'], 0)} €",
          f"{fr(-r['natalite_supplementaire_pct'], 1)} %"] for r in quart],
        [1.7, 1.7, 1.7, 3.1, 1.7, 1.7, 1.8, 2.2], taille=9,
        note="Q1 : le quart des couples au niveau de vie le plus faible ; Q4 : le plus élevé. "
             "Effet estimé : recul de la natalité du quartile que prédit le modèle si l'on "
             "passe du barème de 2014 à celui de 2025.")
    s.puces([
        "**Les classes moyennes ont perdu le plus** : 1 600 à 1 800 € par an en Q2 et Q3, soit "
        "un quart à un tiers de l'incitation, presque entièrement par la PAJE, dont les "
        "plafonds de 2014 et 2018 visent les revenus intermédiaires.",
        "**Les plus modestes sont protégés** : sous les plafonds, ils gardent l'essentiel de "
        "la PAJE (−11 %). **Les plus aisés** n'en touchaient déjà presque plus en 2014 ; ils "
        "perdent surtout par la modulation des allocations familiales de 2015 (−14 %).",
        "**Le quotient familial** recule de 450 € en Q2, non parce que la politique familiale "
        "a été réduite, mais parce que l'impôt a baissé pour tous (première tranche ramenée "
        "de 14 % à 11 % en 2020, décote renforcée) : ces couples paient moins d'impôt, enfant "
        "ou non, mais une demi-part de plus leur rapporte moins. En Q3 et Q4, il progresse "
        "légèrement.",
        "Côté natalité, le modèle prédit un recul de l'ordre de 5 à 8 % dans les classes "
        "moyennes et de 1 à 2 % aux extrémités. La natalité mesurée recule dans tous les "
        "quartiles, sans lien net avec l'ampleur des pertes ; avec 60 à 130 naissances par "
        "quartile et la rupture de mesure de 2022, cette comparaison n'est pas concluante.",
    ])
    s.h2("4.5 L'effet sur la natalité et la part de la baisse")
    s.p("Le gain simulé dépend du revenu, de l'activité et des enfants déjà nés, qui ont leur "
        "propre lien avec la natalité. Pour isoler l'effet du barème, le modèle contrôle le gain "
        "qu'aurait procuré un barème fixe (celui de 2014) : seule reste la variation créée par "
        "les réformes, qui ont touché des familles différentes à des dates différentes. Sans ce "
        "contrôle, l'effet sort même négatif, ce qui montre qu'il est indispensable.")
    s.puces([
        f"Mille euros de gain annuel en plus vont avec une probabilité de naissance plus élevée "
        f"de {fr(pf['ame14'], 2)} point (environ +{fr(pf['ame14_rel'], 0)} %), mais l'effet "
        f"n'est pas significatif (p = {fr(pf['p14'], 2)}).",
        f"Il porte sur les naissances de rang deux ou plus (+{fr(pf['ame14_r2'], 2)} point, "
        f"p = {fr(pf['p14_r2'], 2)}) ; il est nul pour le premier enfant.",
        f"Rendre aux couples de 2025 le barème de 2014 relèverait leur natalité de "
        f"{fr(pf['rel14'], 1)} %, soit {fr(pf['part14'], 0)} % de la baisse de l'ICF. "
        f"L'intervalle de confiance va de {fr(pf['bas14'], 0)} % à +{fr(pf['haut14'], 0)} % : "
        "on ne peut ni exclure un effet nul, ni exclure un effet important.",
    ])
    s.p("Une approche macroéconomique donne un ordre de grandeur plus faible. Selon Saudubray "
        "(2026), 0,5 point de PIB de dépenses familiales en plus va avec 0,05 enfant par femme "
        "à long terme dans les pays de l'OCDE. Les prestations du risque famille sont passées "
        "d'environ 2,4 % du PIB en 2016 à 2,2 % en 2023 : appliqué à ce recul, ce coefficient "
        "donne environ −0,02 enfant par femme, soit 5 % de la baisse.")


def autres_pistes(s):
    s.h1("5. Vie en couple, calendrier, départ de chez les parents")
    s.puces([
        "**Vie en couple** : la part des femmes de 15 à 49 ans vivant en couple passe de 58,3 % "
        "à 53,4 % entre 2014 et 2025, surtout à 25-29 ans (de 65,8 % à 56,1 %). Comme 92 % des "
        "naissances ont lieu en couple, ce recul suffit, à natalité inchangée dans chaque "
        "groupe, à expliquer environ un tiers de la baisse de l'ICF.",
        "**Concentration urbaine** : les jeunes se concentrent dans les grandes agglomérations, "
        "mais la natalité n'y est pas plus faible ; l'effet de composition est nul.",
        "**Calendrier** : l'âge auquel la moitié des couples dispose d'au moins 80 m² recule "
        "d'environ deux ans (de 26,7 à 28,4 ans), mais ce recul coïncide avec la refonte de "
        "l'enquête en 2022. L'âge des mères à la naissance augmente à rythme régulier (de 30,3 "
        "à 31,2 ans selon l'Insee) ; ce report n'explique pas, à lui seul, la baisse.",
        "**Départ de chez les parents** : 36 à 38 % des 18-34 ans vivent chez leurs parents, "
        "sans évolution nette. La moitié est partie vers 22,7 ans pour les femmes et 23,9 ans "
        "pour les hommes, en 2014 comme en 2025, plus tard dans l'agglomération parisienne "
        "(24-25 ans). La colocation n'est pas mesurable dans SRCV.",
    ])


def reste_et_limites(s):
    s.h1("6. Ce qui reste inexpliqué")
    s.p("Environ un tiers de la baisse ne tient à aucun des facteurs mesurés : à âge, vie en "
        "couple, logement et revenu donnés, les couples font moins d'enfants qu'en 2014. Deux "
        "pistes, extérieures à nos données, sont compatibles avec ce constat. Le nombre "
        "d'enfants souhaité recule : chez les femmes de moins de 30 ans, il est passé de 2,5 à "
        "1,9 en vingt ans (Bouchet-Valat et Toulemon, 2025). Et la confiance en l'avenir pèse "
        "durablement sur la fécondité dans les pays de l'OCDE (Saudubray, 2026).")
    s.h1("7. Limites et suites")
    s.puces([
        "Les parts sont des **bornes hautes** : elles lisent comme causales des associations. "
        "Pour le logement, le panel plaide plutôt contre une lecture causale ; pour la "
        "politique familiale, l'estimation est trop imprécise pour trancher.",
        "Le champ du logement et de la politique familiale est celui des **couples constitués** ; "
        "le passage à l'ensemble des femmes suppose que la natalité hors couple varie dans la "
        "même proportion.",
        "Les services (crèches, assistantes maternelles) et le congé parental ne sont pas "
        "observables dans SRCV ; le gain simulé suppose les couples mariés et approche le revenu "
        "imposable.",
        "Pistes : une variation extérieure de l'offre de logement pour trancher la causalité ; "
        "relier la vie en couple et le départ de chez les parents au logement, avec les "
        "fichiers détaillés du recensement ; tester les réformes de la PAJE et des allocations "
        "familiales, rang par rang, sur des données plus larges.",
    ])


def donnees_documents(s):
    s.h1("8. Qualité des données et documents produits")
    s.h2("8.1 Anomalies trouvées dans SRCV")
    s.puces([
        "Les poids des vagues 2022 et 2023 sont anormalement bas pour les ménages avec "
        "nouveau-né, sans poids de remplacement dans les fichiers ; un message à l'ADISP est "
        "prêt. Les statistiques pondérées écartent ces deux vagues.",
        "Depuis 2022, la naissance est reconstruite à partir du fichier individus, et la "
        "moitié des naissances est comptée deux fois d'une vague à l'autre ; le niveau mesuré "
        "chute avec la refonte. Seuls les écarts et les compositions sont comparés dans le "
        "temps, jamais les niveaux.",
        "Corrigés en cours de route : des surfaces aberrantes dans le fichier de panel (des "
        "années saisies comme surfaces), 38 mensualités d'emprunt aberrantes, une variable de "
        "type de logement qui change de nomenclature sous le même nom, et un code de zone "
        "géographique hors nomenclature en 2014 (en fait le Sud-Ouest).",
    ])
    s.h2("8.2 Documents")
    s.tableau(
        ["Document", "Pages", "Contenu"],
        [["Étude SRCV", "42", "Logement, natalité et politique familiale"],
         ["Note de synthèse", "4", "Version courte des résultats sur le logement"],
         ["Note de modélisation", "10", "Modèle à générations imbriquées à trois périodes"],
         ["Article", "20", "« La surface fait-elle les naissances ? », 2014-2025, pour la "
          "Revue économique ou la Revue d'économie politique"],
         ["Page de bilan", "—", "Point d'étape partagé en ligne"]],
        [4.2, 1.6, 10.4])


def sources(s):
    s.h1("Sources")
    for ref in [
        "Insee, enquête Statistiques sur les ressources et conditions de vie (SRCV), vagues "
        "2014 à 2025 ; calculs de l'auteur.",
        "Insee (2026), « Bilan démographique 2025 », Insee Première n° 2087.",
        "Bouchet-Valat M., Toulemon L. (2025), « Les Français·es veulent moins d'enfants », "
        "Population et Sociétés n° 635, Ined.",
        "Saudubray L. (2026), « Les déterminants économiques de la fécondité dans les pays de "
        "l'OCDE », document de travail, Chaire TDTE.",
        "Cnaf, Observatoire national de la petite enfance (Onape), L'accueil du jeune enfant, "
        "éditions 2016 et 2025.",
        "Drees, La protection sociale en France et en Europe (données 2016) ; Insee, "
        "France, portrait social (données 2023).",
        "OpenFisca-France, modèle libre de micro-simulation du système socio-fiscal, utilisé en "
        "local pour les barèmes 2014-2025.",
    ]:
        par = s.doc.add_paragraph()
        par.paragraph_format.space_after = Pt(3)
        texte(par, ref, taille=10)


def chiffres_politique_familiale():
    """Lit les sorties du script 38 sur 2014-2025 (suffixe _2014)."""
    dec = lire("politique_familiale_decomposition_2014.csv")
    mod = lire("politique_familiale_modele_2014.csv")
    ins = lire("politique_familiale_par_instrument_2014.csv")

    def ligne(bareme, ech):
        return next(r for r in dec if r["bareme"] == str(bareme) and r["echantillon"] == ech)

    def modele(ech):
        return next(r for r in mod if r["echantillon"] == ech and "retenu" in r["specification"])

    e = ligne(2014, "Ensemble")
    gain25, gain14 = float(e["gain_2025"]), float(e["gain_bareme"])
    pf = {
        "perte": gain14 - gain25, "perte_pct": 100 * (1 - gain25 / gain14),
        "part14": float(e["part_de_la_baisse_pct"]),
        "bas14": float(e["part_ic_bas"]), "haut14": float(e["part_ic_haut"]),
        "rel14": float(e["ecart_relatif_pct"]),
        "ame14": float(modele("Ensemble")["ame_pts"]), "p14": float(modele("Ensemble")["p.value"]),
        "ame14_r2": float(modele("Rang 2 et plus")["ame_pts"]),
        "p14_r2": float(modele("Rang 2 et plus")["p.value"]),
    }
    pf["ame14_rel"] = 100 * pf["ame14"] / 7.4     # taux de naissance de référence, 7,4 %
    gains = []
    for lib, ech in (("Tous les couples", "Ensemble"), ("Premier enfant", "1er enfant"),
                     ("Rang deux ou plus", "Rang 2 et plus")):
        r = ligne(2014, ech)
        gains.append((lib, {2014: float(r["gain_bareme"]), 2025: float(r["gain_2025"])}))
    libelles = {"Prestations monetaires (AF, CF, PAJE)": "Prestations monétaires",
                "dont PAJE (allocation de base et prime)":
                    "   dont PAJE (allocation de base et prime)",
                "dont allocations familiales et complement familial":
                    "   dont allocations familiales et complément familial",
                "Avantage fiscal (quotient familial)": "Baisse d'impôt (quotient familial)"}
    instr = [(libelles[r["instrument"].strip()],
              {k: float(r[k]) for k in ("gain_bareme", "gain_2025", "ecart_gain",
                                        "part_de_la_baisse_pct", "part_ic_bas", "part_ic_haut")})
             for r in ins if r["bareme"] == "2014"]
    return pf, gains, instr


def main():
    pf, gains, instr = chiffres_politique_familiale()
    quart = [{k: (v if k == "q" else float(v)) for k, v in r.items()}
             for r in lire("politique_familiale_par_quartile_2014.csv")]
    diplomes = [{k: (v if k == "groupe" else float(v)) for k, v in r.items()}
                for r in lire("decomposition_icf_par_diplome.csv")]
    s = Synthese()
    s.titre("Logement, politique familiale et baisse de la fécondité",
            "Synthèse des travaux sur l'enquête SRCV, 2014-2025",
            "Kévin Genna, Institut Louis Bachelier", "Octobre 2026")
    en_bref(s, pf)
    contexte(s)
    vue_ensemble(s, pf, diplomes)
    logement(s)
    politique_familiale(s, pf, gains, instr, quart)
    autres_pistes(s)
    reste_et_limites(s)
    donnees_documents(s)
    sources(s)
    s.enregistrer()


if __name__ == "__main__":
    main()
